"""Word 文档解析器 —— 从 .doc / .docx 中提取文本并解析结构化字段。"""
import logging
import os
import re

logger = logging.getLogger(__name__)

_FORMTEXT_PATTERN = re.compile(r"FORMTEXT[\x01-\x1f]*")


def _clean_formtext(text: str) -> str:
    return _FORMTEXT_PATTERN.sub("", text).strip()


def _extract_first_match(pattern: str, text: str, group: int = 1) -> str:
    m = re.search(pattern, text, re.DOTALL)
    return _clean_formtext(m.group(group)) if m else ""


def _normalize_date_text(value: str) -> str:
    text = str(value or "")
    text = re.sub(r"\s+", "", text)
    text = (
        text.replace("年", "-")
        .replace("月", "-")
        .replace("日", "")
        .replace("/", "-")
        .replace(".", "-")
    )
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if not match:
        return ""
    return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"


def _extract_date(text: str, label: str) -> str:
    pattern = label + r"\s*[：:]?\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日"
    m = re.search(pattern, text)
    if m:
        return f"{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}"
    pattern2 = label + r"\s*[：:]?\s*(\d{4}[年/\-.]\d{1,2}[月/\-.]\d{1,2})"
    m2 = re.search(pattern2, text)
    if m2:
        return _normalize_date_text(m2.group(1))
    return ""


def _extract_date_range(text: str) -> tuple[str, str]:
    range_patterns = [
        r"(?:自|从|起)\s*(\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日)\s*(?:起)?\s*(?:至|到|止)\s*(\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日)",
        r"(?:自|从|起)\s*(\d{4}[年/\-.]\d{1,2}[月/\-.]\d{1,2})\s*(?:起)?\s*(?:至|到|止)\s*(\d{4}[年/\-.]\d{1,2}[月/\-.]\d{1,2})",
        r"有效期(?:自|从|起)\s*(\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日).{0,15}?(?:至|到|止)\s*(\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日)",
        r"服务期(?:限)?(?:自|从)\s*(\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日).{0,15}?(?:至|到|止)\s*(\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日)",
        r"合同期(?:限)?(?:自|从)\s*(\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日).{0,15}?(?:至|到|止)\s*(\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日)",
    ]
    for pattern in range_patterns:
        m = re.search(pattern, text)
        if m:
            start = _normalize_date_text(m.group(1))
            end = _normalize_date_text(m.group(2))
            if start and end:
                return start, end
    return "", ""


def _cleanup_project_name(value: str) -> str:
    text = str(value or "").strip()
    if "考评" in text or "考核" in text:
        return ""
    suffix_match = re.search(
        r"^(.*(?:项目|广场|中心|大厦|园区|小区|天地|综合体|购物广场|商业街))", text
    )
    if suffix_match:
        return suffix_match.group(1).strip(" -_（）()")[:50]
    text = re.sub(
        r"(?:保洁|保安)?(?:外包)?(?:服务|管理)?(?:合同|协议|细则|规则|办法|标准|条款).*$", "", text
    )
    return text.strip(" -_（）()")[:50]


def extract_text_from_doc(file_path: str) -> str:
    import olefile

    ole = olefile.OleFileIO(file_path)
    word_stream = ole.openstream("WordDocument").read()

    text_start = 0
    for offset in range(0x200, min(len(word_stream), 0x4000), 2):
        try:
            ch = word_stream[offset:offset + 2].decode("utf-16-le")
            if "\u4e00" <= ch <= "\u9fff" or ch in "：（）《》、；，。！？":
                chars = []
                j = offset
                while j < len(word_stream) - 2:
                    try:
                        c = word_stream[j:j + 2].decode("utf-16-le")
                        if c == "\x00":
                            break
                        chars.append(c)
                        j += 2
                    except Exception:
                        break
                candidate = "".join(chars)
                if len(candidate) > 50:
                    text_start = offset
                    break
        except Exception:
            continue

    if text_start == 0:
        ole.close()
        return ""

    text_bytes = word_stream[text_start:]
    text = text_bytes.decode("utf-16-le", errors="replace")
    null_pos = text.find("\x00\x00\x00")
    if null_pos > 0:
        text = text[:null_pos]

    cleaned = []
    for ch in text:
        code = ord(ch)
        if code == 0x0D:
            cleaned.append("\n")
        elif code == 0x07:
            cleaned.append("\t")
        elif code == 0x13:
            cleaned.append(" ")
        elif code == 0x14:
            cleaned.append(" ")
        elif code == 0x15:
            cleaned.append(" ")
        elif code <= 0x1F:
            continue
        elif code == 0x00:
            continue
        else:
            cleaned.append(ch)

    result = "".join(cleaned)
    result = re.sub(r"\n{3,}", "\n\n", result)
    result = re.sub(r" {2,}", " ", result)

    ole.close()
    return result


def extract_text_from_docx(file_path: str) -> str:
    from docx import Document

    doc = Document(file_path)
    paragraphs = []

    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            row_texts = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_texts.append(cell_text)
            if row_texts:
                paragraphs.append(" | ".join(row_texts))

    return "\n".join(paragraphs)


def extract_text_from_word(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".doc":
        return extract_text_from_doc(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")


def _extract_first_match_oneline(pattern: str, text: str, group: int = 1) -> str:
    m = re.search(pattern, text)
    return _clean_formtext(m.group(group)) if m else ""


def _parse_metadata(text: str, original_name: str) -> dict:
    text_oneline = text.replace("\n", "")

    contract_no = _extract_first_match_oneline(
        r"合同编号[：:]\s*([^\n\d一二三四五六七八九十第条章]{2,30}?)(?:[\s\n]|第|$)",
        text,
    )
    if not contract_no:
        contract_no = _extract_first_match_oneline(
            r"合同编号[：:]\s*([^\n]{2,30})", text
        )

    supplier = _extract_first_match_oneline(
        r"乙\s*方[：:]\s*([^\n]*)", text
    )
    _noise_keywords = ["盖章", "地址", "地 址", "联系人", "电话", "开户行", "法定代表",
                       "授权代表", "邮编", "邮箱", "传真", "日期", "签名", "签字"]
    if any(kw in supplier for kw in _noise_keywords):
        supplier = ""

    project_name = _extract_first_match_oneline(
        r"保洁区域[：:]?\s*([^\n；;。]{2,50})", text
    )
    if not project_name:
        project_name = _extract_first_match_oneline(
            r"服务项目[：:]?\s*([^\n；;。]{2,50})", text
        )
    if not project_name:
        project_name = _extract_first_match_oneline(
            r"项目名称[：:]?\s*([^\n；;。]{2,50})", text
        )
    if not project_name:
        project_name = _extract_first_match_oneline(
            r"项目(?!(?:名称|编号|周期|性质|属性|概况|简介|负责|经理|总监|地址|地点|位置|类型|内容|说明|范围|概述|代码|类别|分类|管理|主管|所在))[：:]?\s+([^\n；;。]{2,50})", text
        )
    if project_name:
        project_name = _cleanup_project_name(project_name)

    contract_name = os.path.splitext(original_name)[0]

    service_type = ""
    if "保洁" in text_oneline[:200]:
        service_type = "保洁"
    elif "保安" in text_oneline[:200]:
        service_type = "保安"

    business_type = ""
    business_keywords = ["住宅", "商业", "物业", "酒店", "街区", "写字楼", "保洁", "保安"]
    for kw in business_keywords:
        if kw in text_oneline[:500]:
            business_type = kw
            break

    start_date, end_date = _extract_date_range(text)

    if not start_date:
        start_date = _extract_date(text, "自")
    if not end_date:
        end_date = _extract_date(text, "至")

    if not start_date:
        for pat in [
            r"(?:开始日期|合同开始|合同期自|服务期自|起止日期自|起始日期)\s*[：:]?\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日",
            r"(?:开始日期|合同开始|合同期自|服务期自|起止日期自|起始日期)\s*[：:]?\s*(\d{4}[年/\-.]\d{1,2}[月/\-.]\d{1,2})",
            r"合同有效期[^自]*(?:自|从)\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日",
        ]:
            m = re.search(pat, text)
            if m:
                if m.lastindex == 3:
                    start_date = f"{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}"
                else:
                    start_date = _normalize_date_text(m.group(1))
                break

    if not end_date:
        for pat in [
            r"(?:结束日期|合同结束|合同期至|服务期至|起止日期至|终止日期)\s*[：:]?\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日",
            r"(?:结束日期|合同结束|合同期至|服务期至|起止日期至|终止日期)\s*[：:]?\s*(\d{4}[年/\-.]\d{1,2}[月/\-.]\d{1,2})",
            r"合同有效期.*?(?:至|到)\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日",
        ]:
            m = re.search(pat, text)
            if m:
                if m.lastindex == 3:
                    end_date = f"{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}"
                else:
                    end_date = _normalize_date_text(m.group(1))
                break

    return {
        "project_name": project_name,
        "business_type": business_type,
        "supplier": supplier,
        "contract_no": contract_no,
        "contract_name": contract_name,
        "service_type": service_type,
        "start_date": start_date,
        "end_date": end_date,
    }


_DEFAULT_RULES = [
    {
        "code": "S01", "name": "实际到岗数", "penalty_type": "formula",
        "condition": "低于合同约定岗位数",
        "penalty_formula": "工时单价 * 1.2 * 缺勤总时长",
        "enabled": True,
    },
    {
        "code": "S02", "name": "未在约定时间调换人员", "penalty_type": "formula",
        "condition": "对于不适合本项目甲方要求调换的保洁人员，乙方必须在10天内更换完成，否则自第11天起",
        "penalty_formula": "工时单价 * 0.1 * 在岗总时长",
        "enabled": True,
    },
    {
        "code": "S03", "name": "超龄人员", "penalty_type": "formula",
        "condition": "超过合同约定年龄，超龄总人数≤当日在岗总人数的20%不扣款",
        "penalty_formula": "工时单价 * 0.1 * 在岗总时长",
        "enabled": True,
    },
    {
        "code": "S04", "name": "考勤管理", "penalty_type": "composite",
        "condition": "漏打卡、迟到、早退、脱岗等",
        "sub_rules": [
            {"code": "S04-1", "name": "漏打卡", "amount": 50, "unit": "元/人次", "free_limit": 3, "free_unit": "次/月"},
            {"code": "S04-2", "name": "迟到或早退在30分钟以内", "amount": 20, "unit": "元/次"},
            {"code": "S04-3", "name": "迟到或早退在60分钟以内", "amount": 50, "unit": "元/次"},
            {"code": "S04-4", "name": "脱岗", "amount": "当日服务费*2.5", "unit": "元"},
        ],
        "enabled": True,
    },
    {
        "code": "S05", "name": "员工形象及行为规范", "penalty_type": "composite",
        "condition": "乙方员工的仪容仪表及行为规范严重影响甲方形象或屡教不改者",
        "sub_rules": [
            {"code": "S05-1", "name": "一般违规", "amount": 100, "unit": "元/人/次"},
            {"code": "S05-2", "name": "吸烟/睡岗", "amount": 200, "unit": "元/人/次"},
        ],
        "enabled": True,
    },
    {
        "code": "S06", "name": "不合格项整改", "penalty_type": "fixed",
        "condition": "相同的整改内容一个月中累计出现三次及以上",
        "amount": 100, "unit": "元/张",
        "enabled": True,
    },
    {
        "code": "S07", "name": "随意倾倒垃圾", "penalty_type": "fixed",
        "condition": "乙方人员将垃圾、杂物等倒入甲方排水管道",
        "amount": 500, "unit": "元/次",
        "enabled": True,
    },
    {
        "code": "S08", "name": "盗窃等违法行为", "penalty_type": "fixed",
        "condition": "乙方人员在工作期间发生盗窃等违法行为",
        "amount": 2000, "unit": "元",
        "enabled": True,
    },
    {
        "code": "S09", "name": "闹事", "penalty_type": "percent",
        "condition": "乙方或乙方员工发生闹事行为",
        "tiers": [
            {"times": 1, "percent": 15, "base": "当月服务费"},
            {"times": 2, "percent": 30, "base": "当月服务费"},
        ],
        "enabled": True,
    },
    {
        "code": "S10", "name": "员工休息管理", "penalty_type": "fixed",
        "condition": "未在休息室休息，在场内乱窜等",
        "amount": 50, "unit": "元/次",
        "enabled": True,
    },
    {
        "code": "S11", "name": "行为规范", "penalty_type": "composite",
        "condition": "按甲方《行为规范 物业现场员工行为规范》执行",
        "sub_rules": [
            {"code": "S11-1", "name": "一般违规", "amount": 50, "unit": "元/人次"},
            {"code": "S11-2", "name": "随地吐痰/大小便", "amount": 200, "unit": "元/人次"},
        ],
        "enabled": True,
    },
    {
        "code": "S12", "name": "人员流动率", "penalty_type": "formula",
        "condition": "保洁人员整体流动率≥20%",
        "penalty_formula": "总扣分数 * 人均服务费",
        "enabled": True,
    },
    {
        "code": "S13", "name": "保险真伪", "penalty_type": "fixed",
        "condition": "恶意造假",
        "amount": 5000, "unit": "元/次",
        "enabled": True,
    },
]


def _parse_deduction_rules(text: str) -> list:
    attachment_text = _extract_attachment3(text)
    if not attachment_text:
        return _DEFAULT_RULES

    rules = _parse_rules_from_table(attachment_text)
    if not rules or len(rules) < 5:
        return _DEFAULT_RULES

    return rules


def _extract_attachment3(text: str) -> str:
    patterns = [
        r"附件\s*3\s*\n+外包保洁服务考核表",
        r"序号\s+检查内容\s+评判标准\s+不符合项扣款金额",
    ]
    for p in patterns:
        m = re.search(p, text)
        if m:
            start = m.start()
            end = text.find("附件4", start)
            if end == -1:
                end = text.find("备注：", start)
                if end == -1:
                    end = min(start + 5000, len(text))
            return text[start:end]
    return ""


def _parse_rules_from_table(attachment_text: str) -> list:
    items = re.split(r"\n\s*(?=\d{1,2}\s{2,})", attachment_text)
    rules = []
    for item in items:
        rule = _parse_single_rule_item(item)
        if rule:
            rules.append(rule)
    return rules


def _parse_single_rule_item(item_text: str) -> dict | None:
    m = re.match(r"(\d{1,2})\s{2,}", item_text)
    if not m:
        return None
    seq = int(m.group(1))
    if seq < 1 or seq > 13:
        return None

    content = item_text[m.end():].strip()
    template = _DEFAULT_RULES[seq - 1] if seq <= len(_DEFAULT_RULES) else None
    if not template:
        return None

    if template["penalty_type"] == "formula":
        formula = _extract_formula(content)
        return {**template, "penalty_formula": formula or template["penalty_formula"]}

    if template["penalty_type"] == "fixed":
        amounts = re.findall(r"(\d+)\s*元", content)
        return {**template, "amount": int(amounts[0]) if amounts else template["amount"]}

    if template["penalty_type"] == "composite":
        sub_rules = _extract_sub_rules(content, template.get("sub_rules", []))
        return {**template, "sub_rules": sub_rules}

    if template["penalty_type"] == "percent":
        return {**template}

    return template


def _extract_formula(text: str) -> str:
    m = re.search(r"工时单价\s*\*\s*[\d.]+", text)
    if m:
        return m.group(0)
    return ""


def _extract_sub_rules(text: str, template_sub: list) -> list:
    amounts = re.findall(r"(\d+)\s*元", text)
    if not amounts:
        return template_sub
    result = []
    for i, sub in enumerate(template_sub):
        if i < len(amounts) and isinstance(sub.get("amount"), int):
            amount = int(amounts[i])
            if amount <= 10000:
                result.append({**sub, "amount": amount})
            else:
                result.append(sub)
        else:
            result.append(sub)
    return result


def parse_contract(file_path: str, original_name: str, ext: str) -> dict:
    text = extract_text_from_word(file_path)
    if not text:
        logger.warning(f"无法从文件 {original_name} 中提取文本")
        return {
            "project_name": "",
            "business_type": "",
            "supplier": "",
            "contract_no": "",
            "contract_name": os.path.splitext(original_name)[0],
            "service_type": "",
            "start_date": "",
            "end_date": "",
            "rules": [],
        }

    metadata = _parse_metadata(text, original_name)
    rules = _parse_deduction_rules(text)

    return {
        **metadata,
        "rules": rules,
    }