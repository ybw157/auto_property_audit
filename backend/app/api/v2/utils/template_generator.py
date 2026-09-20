from pathlib import Path
from urllib.parse import quote
from fastapi.responses import FileResponse
from openpyxl import Workbook
from app.api.v2.core.config import settings


def _disposition(filename: str) -> dict:
    """构造 Content-Disposition 头，同时提供 ASCII fallback 和 UTF-8 编码名。"""
    ascii_name = filename.encode("ascii", "replace").decode("ascii")
    utf8_name = quote(filename)
    return {
        "Content-Disposition": (
            f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{utf8_name}'
        ),
    }


def generate_project_template():
    """生成「项目基础资料模板.xlsx」。

    结构与解析器（excel_project_parser）期望的表头布局完全一致，确保下载后按规范
    填写、上传即可被正确解析：

      - 基础信息：表头行（业态 / 外包公司 / 合同约定岗位数 / 实际在岗岗位数）。
        解析器只按表头关键词识别这几列，不读「项目名称」（项目名一律由登录账号决定，
        表格里的项目名不解析、不读取），因此模板不再提供「项目名称」填写列。
      - 合同编制表：表头行（业态 / 岗位名称 / 岗位归属 / 责任区域 / 人员姓名 /
        岗位时间 / 日时长 / 工时单价 / 实际岗位数 / 月度合价），与 _detect_contract_columns 对应。
      - 月度排班表：表头行（业态 / 岗位名称 / 固定区域 + 1..31 日），与 parse_schedule_sheet 对应。
      - 班次标准：仅作人工参考，不参与解析。

    每次下载都重新生成（不再缓存旧文件），保证模板始终与最新解析规则同步。
    """
    path = Path(settings.template_dir) / "项目基础资料模板.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "基础信息"
    ws.append(["业态", "外包公司", "合同约定岗位数", "实际在岗岗位数"])

    ws = wb.create_sheet("合同编制表")
    ws.append([
        "业态", "岗位名称", "岗位归属", "责任区域", "人员姓名",
        "岗位时间", "日时长", "工时单价", "实际岗位数", "月度合价",
    ])

    ws = wb.create_sheet("月度排班表")
    schedule_header = ["业态", "岗位名称", "固定区域"] + [str(d) for d in range(1, 32)]
    ws.append(schedule_header)

    ws = wb.create_sheet("班次标准")
    ws.append(["班次名称", "上班时间", "下班时间", "工时", "要求打卡次数", "允许打卡窗口"])

    wb.save(path)
    return FileResponse(path, headers=_disposition("项目基础资料模板.xlsx"))


def generate_attendance_template():
    path = Path(settings.template_dir) / "BI考勤模板.xlsx"
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "BI考勤"
        ws.append(["日期", "姓名", "打卡时间"])
        wb.save(path)
    return FileResponse(path, headers=_disposition("BI考勤模板.xlsx"))
