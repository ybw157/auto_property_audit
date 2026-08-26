from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from app.core.config import settings

def generate_excel_report(batch_id: int, results: dict) -> dict:
    report_dir = Path(settings.report_dir) / str(batch_id)
    report_dir.mkdir(parents=True, exist_ok=True)
    file_path = report_dir / "物业外包智能审核结果.xlsx"
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    add_attendance_sheet(wb, "保洁考勤明细", results.get("attendance_details", []))
    add_sheet(wb, "排班审核汇总", results.get("position_fulfillment", []))
    add_sheet(wb, "人员异常统计", results.get("exception_statistics", []))
    add_sheet(wb, "人员考勤扣款", results.get("attendance_deductions", []))
    add_sheet(wb, "扣款汇总", results.get("deduction_summary", []))

    wb.save(file_path)
    return {"report_type": "综合审核结果", "file_format": "xlsx", "file_path": str(file_path), "download_url": f"/files/reports/{batch_id}/物业外包智能审核结果.xlsx"}

def add_attendance_sheet(wb, title: str, rows: list[dict]):
    """保洁考勤明细专用：正常行黑字白底，异常行红字。"""
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["暂无数据"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="EDE9FE")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    red_font = Font(color="FF0000")
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
        exception_type = str(row.get("exception_type", "") or "")
        deduction_amount = row.get("deduction_amount", 0) or 0
        if exception_type or float(deduction_amount) > 0:
            for cell in ws[ws.max_row]:
                cell.font = red_font
    ws.freeze_panes = "A2"
    for col_idx, header in enumerate(headers, start=1):
        width = max(len(str(header)) + 4, 14)
        for row in rows[:50]:
            width = max(width, min(len(str(row.get(header, ""))) + 2, 36))
        ws.column_dimensions[get_column_letter(col_idx)].width = width

def add_sheet(wb, title: str, rows: list[dict]):
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["暂无数据"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="EDE9FE")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.freeze_panes = "A2"
    for col_idx, header in enumerate(headers, start=1):
        width = max(len(str(header)) + 4, 14)
        for row in rows[:50]:
            width = max(width, min(len(str(row.get(header, ""))) + 2, 36))
        ws.column_dimensions[get_column_letter(col_idx)].width = width
