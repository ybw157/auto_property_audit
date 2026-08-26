from datetime import datetime, date, time
from pathlib import Path
from typing import Any
import re
import csv
from openpyxl import load_workbook

from app.services.bi_parser import resolve_bi_path, parse_daily_clock

PROJECT_INFO_ALIASES = ["项目基础信息", "项目基础表", "基础信息", "Sheet1项目基础信息", "Sheet1"]
CONTRACT_ALIASES = ["合同编制表", "保洁人员实际在岗编制表", "合同编制", "Sheet2合同编制表", "Sheet2"]
SCHEDULE_ALIASES = ["月度排班表", "排班表", "Sheet3月度排班表", "Sheet3"]
SHIFT_ALIASES = ["班次标准"]
REQUIRED_SHEET_GROUPS = [
    ("项目基础信息", PROJECT_INFO_ALIASES),
    ("合同编制表", CONTRACT_ALIASES),
    ("月度排班表", SCHEDULE_ALIASES),
]

def validate_project_workbook(path: str | Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    missing = [label for label, aliases in REQUIRED_SHEET_GROUPS if not find_sheet_name(wb, aliases)]
    if missing:
        return {"valid": False, "missing_sheets": missing, "message": f"缺少Sheet：{','.join(missing)}"}
    return {"valid": True, "missing_sheets": [], "message": "模板校验通过"}

def parse_project_workbook(path: str | Path) -> dict:
    wb = load_workbook(path, data_only=True)
    project_ws = wb[find_sheet_name(wb, PROJECT_INFO_ALIASES)]
    contract_ws = wb[find_sheet_name(wb, CONTRACT_ALIASES)]
    schedule_ws = wb[find_sheet_name(wb, SCHEDULE_ALIASES)]
    shift_name = find_sheet_name(wb, SHIFT_ALIASES)

    project_info = parse_project_info(project_ws)
    if not project_info.get("审核月份"):
        project_info["审核月份"] = infer_month_from_filename(path)
    contracts = parse_contract_staffing(contract_ws)
    schedules = parse_monthly_schedule(schedule_ws, project_info.get("审核月份", infer_month_from_filename(path)))
    shifts = parse_table_sheet(wb[shift_name]) if shift_name else extract_shift_rules_from_contracts(contracts)
    return {
        "project_info": project_info,
        "contract_staffing": contracts,
        "monthly_schedules": schedules,
        "shift_standards": shifts,
    }

def parse_attendance_workbook(path: str | Path | list[str] | tuple[str, ...]) -> list[dict]:
    if isinstance(path, (list, tuple)):
        rows = []
        for item in path:
            rows.extend(parse_attendance_workbook(item))
        return rows
    path = Path(path)
    if path.suffix.lower() == ".csv":
        return parse_attendance_table(read_csv_rows(path), path)
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [[clean_cell(cell.value) for cell in row] for row in ws.iter_rows()]
    return parse_attendance_table(rows, path)

def read_csv_rows(path: Path) -> list[list[Any]]:
    last_error = None
    for encoding in ["utf-8-sig", "gb18030", "gbk"]:
        try:
            with path.open("r", encoding=encoding, newline="") as f:
                return [[clean_cell(value) for value in row] for row in csv.reader(f)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"无法识别CSV编码：{path.name}，{last_error}")

def parse_attendance_table(rows: list[list[Any]], path: Path) -> list[dict]:
    if not rows:
        return []
    header_index = find_attendance_header_index(rows)
    if header_index is None:
        return []
    headers = [str(v).strip() for v in rows[header_index]]
    if has_day_columns(headers):
        return parse_wide_attendance_rows(headers, rows[header_index + 1 :], path)
    table_rows = []
    for row_no, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not any(v not in ("", None) for v in row):
            continue
        item = {headers[i] if i < len(headers) and headers[i] else f"空列{i+1}": row[i] if i < len(row) else "" for i in range(len(headers))}
        item["_source_row_no"] = row_no
        table_rows.append(item)
    return table_rows

def find_attendance_header_index(rows: list[list[Any]]) -> int | None:
    for idx, row in enumerate(rows[:10]):
        values = [str(v).strip() for v in row]
        has_name = "姓名" in values or "人员姓名" in values or "员工姓名" in values
        has_month = "月份" in values or "日期" in values or "考勤日期" in values
        has_day = any(re.fullmatch(r"\d{1,2}日?", value) for value in values)
        if has_name and (has_month or has_day):
            return idx
    return None

def has_day_columns(headers: list[str]) -> bool:
    return any(re.fullmatch(r"\d{1,2}日?", h) for h in headers)

def parse_wide_attendance_rows(headers: list[str], data_rows: list[list[Any]], path: Path) -> list[dict]:
    day_columns = []
    for idx, header in enumerate(headers):
        match = re.fullmatch(r"(\d{1,2})日?", str(header).strip())
        if match:
            day_columns.append((idx, int(match.group(1))))
    name_idx = find_header_index(headers, ["姓名", "人员姓名", "员工姓名"])
    month_idx = find_header_index(headers, ["月份", "考勤月份"])
    project_idx = find_header_index(headers, ["项目", "项目名称"])
    department_idx = find_header_index(headers, ["部门"])
    if name_idx is None:
        return []
    # 没有月份列时，从文件名推断月份
    inferred_month = infer_month_from_filename(path) if month_idx is None else ""
    result = []
    for row_no, row in enumerate(data_rows, start=2):
        employee = str(row[name_idx] if name_idx < len(row) else "").strip()
        if not employee:
            continue
        if month_idx is not None:
            month_value = row[month_idx] if month_idx < len(row) else ""
            if not month_value:
                continue
            month_text = normalize_attendance_month(month_value, path)
        else:
            month_text = inferred_month
            if not month_text:
                continue
        bi_project = row[project_idx] if project_idx is not None and project_idx < len(row) else ""
        department = row[department_idx] if department_idx is not None and department_idx < len(row) else ""

        # 解析BI路径 -> 标准项目名、业态、服务类型
        bi_meta = resolve_bi_path(str(bi_project), str(department)) if bi_project else {}

        for col_idx, day in day_columns:
            clock_value = row[col_idx] if col_idx < len(row) else ""
            if clock_value in ("", None):
                continue

            # 解析每日打卡时间
            clock_info = parse_daily_clock(clock_value)

            result.append({
                "日期": f"{month_text}-{day:02d}",
                "姓名": employee,
                "打卡时间": clock_value,
                "打卡明细": clock_info.get("clock_times", []),
                "上班打卡": clock_info.get("clock_in"),
                "下班打卡": clock_info.get("clock_out"),
                "实际工时": clock_info.get("work_hours", 0.0),
                "项目": bi_meta.get("project_name") or bi_project,
                "BI原始项目": bi_project,
                "部门": department,
                "服务类型": bi_meta.get("service_type") or department,
                "业态": bi_meta.get("business_type") or infer_business_type_from_attendance(path.name, employee, bi_project, department),
                "_source_file": path.name,
                "_source_row_no": row_no,
            })
    return result

def infer_business_type_from_attendance(file_name: str, employee: str = "", project: str = "", department: str = "") -> str:
    text = f"{file_name} {employee} {project} {department}"
    if "住宅" in text:
        return "住宅"
    if "外场" in text or "外围" in text or "停车场" in text:
        return "外场"
    if "内场" in text or "商场" in text or "商业" in text:
        return "商场"
    return ""

def find_header_index(headers: list[str], names: list[str]) -> int | None:
    for name in names:
        if name in headers:
            return headers.index(name)
    return None

def normalize_attendance_month(value: Any, path: Path) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    text = str(value).strip()
    match = re.fullmatch(r"(20\d{2})(\d{2})", text)
    if match:
        return f"{match.group(1)}-{match.group(2)}"
    match = re.search(r"(20\d{2})[-年/]?(\d{1,2})", text)
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}"
    return infer_month_from_filename(path)

def find_sheet_name(wb, aliases: list[str]) -> str | None:
    alias_set = {normalize_sheet_name(name) for name in aliases}
    for name in aliases:
        if name in wb.sheetnames:
            return name
    for sheet_name in wb.sheetnames:
        normalized = normalize_sheet_name(sheet_name)
        if normalized in alias_set:
            return sheet_name
    for sheet_name in wb.sheetnames:
        normalized = normalize_sheet_name(sheet_name)
        if any(alias and alias in normalized for alias in alias_set):
            return sheet_name
    # 支持"x月排班表"等带月份前缀的排班表sheet名
    if any("排班表" in alias for alias in alias_set):
        for sheet_name in wb.sheetnames:
            normalized = normalize_sheet_name(sheet_name)
            if re.search(r"\d*月?排班表", normalized):
                return sheet_name
    return None

def normalize_sheet_name(value: str) -> str:
    text = str(value or "").strip()
    text = text.replace("（", "(").replace("）", ")")
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"^\d+[.\-_、]*", "", text)
    text = re.sub(r"^sheet\d+[.\-_、:：]*", "sheet", text, flags=re.IGNORECASE)
    text = re.sub(r"\([^)]*\)$", "", text)
    return text.lower()

def infer_month_from_filename(path: str | Path) -> str:
    name = Path(path).name
    match = re.search(r"(20\d{2})[-年._]?(\d{1,2})", name)
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}"
    match = re.search(r"(\d{1,2})[.月]", name)
    if match:
        return f"{datetime.now().year}-{int(match.group(1)):02d}"
    return datetime.now().strftime("%Y-%m")

def parse_project_info(ws) -> dict:
    headers = [clean_cell(c.value) for c in ws[1]]
    row = [clean_cell(c.value) for c in ws[2]] if ws.max_row >= 2 else []
    data = {str(headers[i]): row[i] if i < len(row) else "" for i in range(len(headers)) if headers[i] not in ("", None)}
    return {
        "项目名称": data.get("项目名称", ""),
        "项目编码": data.get("项目编码", data.get("项目名称", "")),
        "业态": data.get("业态", ""),
        "供应商": data.get("保洁外包公司") or data.get("保安外包公司") or data.get("供应商", ""),
        "合同编号": data.get("合同编号", ""),
        "审核月份": data.get("审核月份", ""),
        "负责人": data.get("负责人", ""),
        "模板版本": data.get("模板版本", "江宁项目模板"),
        "备注": data.get("备注", ""),
    }

def parse_contract_staffing(ws) -> list[dict]:
    rows = parse_table_sheet(ws)
    result = []
    last_project = ""
    last_business = ""
    for row in rows:
        if row.get("项目名称"):
            last_project = row.get("项目名称")
        if row.get("业态"):
            last_business = row.get("业态")
        position = str(row.get("岗位名称") or row.get("岗位") or "").strip()
        if not position:
            continue
        headcount = row.get("实际岗位数") or row.get("合同人数") or 0
        result.append({
            "项目名称": last_project,
            "业态": last_business,
            "岗位": position,
            "区域": row.get("岗位责任区域", row.get("区域", "")),
            "合同人数": headcount,
            "编制人员": row.get("岗位人员姓名（考勤参考）", row.get("编制人员", "")),
            "工时单价": row.get("工时单价（元）", row.get("工时单价", 0)),
            "服务类型": row.get("岗位归属", row.get("服务类型", "保洁")),
            "岗位时间": row.get("岗位时间", ""),
            "单岗日时长（h）": row.get("单岗日时长（h）", 8),
            "是否跨天": row.get("是否跨天", "否"),
        })
    return result

def parse_monthly_schedule(ws, audit_month: str) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [clean_cell(v) for v in rows[0]]
    business_idx = find_header_index(headers, ["业态", "业务类型", "服务类型"])
    position_idx = find_header_index(headers, ["岗位名称", "岗位"])
    area_idx = find_header_index(headers, ["区域", "岗位责任区域", "服务区域"])
    shift_idx = find_header_index(headers, ["班次", "班次名称", "默认班次", "岗位时间"])
    if position_idx is None:
        position_idx = 1
    if area_idx is None:
        area_idx = 2
    if business_idx is None and position_idx != 0:
        business_idx = 0
    day_columns = []
    for idx, header in enumerate(headers):
        text = str(header).strip()
        if text.isdigit():
            day_columns.append((idx, int(text)))
    result = []
    current_business = ""
    for source_row_no, row in enumerate(rows[1:], start=2):
        values = [clean_cell(v) for v in row]
        business = values[business_idx] if business_idx is not None and business_idx < len(values) and values[business_idx] not in ("", None) else current_business
        if business:
            current_business = business
        position = str(values[position_idx] if position_idx < len(values) else "").strip()
        area = values[area_idx] if area_idx < len(values) else ""
        shift_name = str(values[shift_idx] if shift_idx is not None and shift_idx < len(values) else "").strip() or position
        if not position:
            continue
        for col_idx, day in day_columns:
            cell_value = values[col_idx] if col_idx < len(values) else ""
            for employee in split_employee_cell(cell_value):
                result.append({
                    "日期": f"{audit_month}-{day:02d}",
                    "岗位": position,
                    "区域": area,
                    "姓名": employee,
                    "班次": shift_name,
                    "排班类型": "正常",
                    "原岗位": "",
                    "原人员": "",
                    "备注": f"业态：{business}" if business else "",
                    "_source_row_no": source_row_no,
                })
    return result

def split_employee_cell(value: Any) -> list[str]:
    if value in (None, "", 0, "0"):
        return []
    text = str(value).strip()
    if not text or text in {"休息", "休", "调休", "无", "-", "/"}:
        return []
    for sep in ["\n", "、", "，", ",", ";", "；"]:
        text = text.replace(sep, " ")
    return [part.strip() for part in text.split() if part.strip() and part.strip() not in {"休息", "休", "0"}]

def extract_shift_rules_from_contracts(contracts: list[dict]) -> list[dict]:
    result = []
    for row in contracts:
        position = row.get("岗位", "")
        if not position:
            continue
        start, end = split_position_time(row.get("岗位时间", ""))
        result.append({
            "班次名称": position,
            "上班时间": start,
            "下班时间": end,
            "工时": row.get("单岗日时长（h）", 8) or 8,
            "要求打卡次数": 2,
            "允许打卡窗口": "",
            "规则来源": "保洁人员实际在岗编制表",
        })
    return result

def split_position_time(value: Any) -> tuple[str, str]:
    text = str(value or "").replace("：", ":").strip()
    if "-" in text:
        left, right = text.split("-", 1)
        return left.strip(), right.strip()
    return "08:00", "16:00"

def parse_key_value_sheet(ws) -> dict:
    data = {}
    for row in ws.iter_rows(values_only=True):
        cells = [clean_cell(v) for v in row]
        if len(cells) >= 2 and cells[0]:
            data[str(cells[0])] = cells[1]
    return data

def parse_table_sheet(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    header_index = None
    headers = []
    for idx, row in enumerate(rows[:10]):
        values = [clean_cell(v) for v in row]
        non_empty = [v for v in values if v not in ("", None)]
        if len(non_empty) >= 2:
            header_index = idx
            headers = [str(v).strip() if v not in (None, "") else f"空列{col+1}" for col, v in enumerate(values)]
            break
    if header_index is None:
        return []
    data = []
    for row_no, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = [clean_cell(v) for v in row]
        if not any(v not in ("", None) for v in values):
            continue
        item = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
        item["_source_row_no"] = row_no
        data.append(item)
    return data

def clean_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, str):
        return value.strip()
    return value
